"""
NeuroStack - Knowledge Service
Service for retrieving relevant knowledge from the vector database (RAG).

This module provides the interface for retrieving relevant knowledge
from the vector database. The actual implementation will be completed
when the Knowledge Service is implemented.
"""

import asyncio
import io
import re
import time
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Final, Literal, Optional
from uuid import UUID

from src.config.logging import get_logger
from src.ctx.models import KnowledgeChunk

logger = get_logger(__name__)


# ============================================
# DOCX / textual PDF extraction (neurostack-credit-system Group 4a)
# ============================================
#
# Parser safety (``pypdf`` and ``python-docx``/``lxml``, pinned in
# ``poetry.lock``).  Both are pure parsers: no network, no macro or script
# execution.  The one external program ``pypdf`` may start -- the
# ``jbig2dec`` binary, for a JBIG2-encoded stream -- is switched off for every
# extraction (``jbig2dec_binary=None``).  Resources are bounded, whatever the
# enforcement: ``pypdf``'s own per-stream output limits (75 MB each), a page
# cap, the declared uncompressed size and entry count of a DOCX archive
# (``python-docx`` reads every part into memory; ``zipfile`` never inflates
# an entry beyond its declared size), and one deadline over all the files of
# a request, checked between pages and files.  The extraction runs in a
# worker thread (:func:`extract_documents`), never on the event loop.  A file
# beyond a bound is unreadable (``422``).

KnowledgeFileKind = Literal["pdf", "docx"]
# Stable 422 codes of ``POST /v1/knowledge/documents/upload``.
KNOWLEDGE_FILE_UNSUPPORTED = "knowledge_file_unsupported"
KNOWLEDGE_FILE_UNREADABLE = "knowledge_file_unreadable"
# Technical bounds of the extraction (not the PDF's commercial limits).
KNOWLEDGE_EXTRACTION_SECONDS: Final = 30.0
KNOWLEDGE_PDF_MAX_PAGES: Final = 5_000
KNOWLEDGE_DOCX_MAX_UNCOMPRESSED_BYTES: Final = 100_000_000
KNOWLEDGE_DOCX_MAX_ENTRIES: Final = 5_000
_PDF_SIGNATURE = b"%PDF-"
_ZIP_SIGNATURE = b"PK\x03\x04"
_DOCX_MAIN_PART = "word/document.xml"
_UNPAIRED_SURROGATE = re.compile("[\ud800-\udfff]")


class KnowledgeFileError(ValueError):
    """A file the upload cannot ingest; ``code`` is the stable answer."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class ExtractedDocument:
    """The text of one uploaded file and what the PDF's rules look at."""

    kind: KnowledgeFileKind
    text: str
    size_bytes: int
    encrypted: bool


def _storable(text: str) -> str:
    """PostgreSQL ``text`` holds neither NUL nor an unpaired surrogate."""

    return _UNPAIRED_SURROGATE.sub("\ufffd", text.replace("\x00", ""))


def _unreadable() -> KnowledgeFileError:
    return KnowledgeFileError(KNOWLEDGE_FILE_UNREADABLE)


def _check_deadline(deadline: float | None) -> None:
    if deadline is not None and time.monotonic() > deadline:
        raise _unreadable()


def _pdf_text(data: bytes, deadline: float | None) -> tuple[str, bool]:
    from pypdf import PdfReader, apply_configuration

    try:
        # Never start the jbig2dec program on an uploaded stream (and let no
        # legacy module constant of pypdf turn it back on for this reader).
        with apply_configuration(jbig2dec_binary=None, disable_legacy_handling=True):
            reader = PdfReader(io.BytesIO(data), strict=False)
            encrypted = bool(reader.is_encrypted)
            if encrypted and not reader.decrypt(""):
                # A password is needed: nothing can be read.
                raise _unreadable()
            if len(reader.pages) > KNOWLEDGE_PDF_MAX_PAGES:
                raise _unreadable()
            pages: list[str] = []
            for page in reader.pages:
                _check_deadline(deadline)
                pages.append(page.extract_text() or "")
    except KnowledgeFileError:
        raise
    except Exception as error:  # a malformed PDF raises many kinds
        raise _unreadable() from error
    return "\n".join(pages), encrypted


def _docx_text(data: bytes, deadline: float | None) -> str:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if _DOCX_MAIN_PART not in {entry.filename for entry in entries}:
                raise KnowledgeFileError(KNOWLEDGE_FILE_UNSUPPORTED)
            # python-docx loads every part: refuse an archive that declares
            # more than the bound (an entry never inflates beyond it).
            if (
                len(entries) > KNOWLEDGE_DOCX_MAX_ENTRIES
                or sum(entry.file_size for entry in entries)
                > KNOWLEDGE_DOCX_MAX_UNCOMPRESSED_BYTES
            ):
                raise _unreadable()
        _check_deadline(deadline)
        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    except KnowledgeFileError:
        raise
    except Exception as error:  # python-docx/lxml raise many kinds
        raise _unreadable() from error
    return "\n".join(parts)


def extract_document_text(
    data: bytes, *, deadline: float | None = None
) -> ExtractedDocument:
    """Text of a DOCX or textual PDF, by content signature (never by the
    declared type).  Unsupported formats, corrupt files, a PDF that needs a
    password, a document with no extractable text (an OCR-only PDF, owner
    decision H) and a file beyond a technical bound (``deadline``: a
    ``time.monotonic()`` instant) raise :class:`KnowledgeFileError` at any
    time; the PDF's commercial limits are the caller's, after the
    enforcement.  Blocking: call it through :func:`extract_documents`."""

    if data.startswith(_PDF_SIGNATURE):
        raw, encrypted = _pdf_text(data, deadline)
        kind: KnowledgeFileKind = "pdf"
    elif data.startswith(_ZIP_SIGNATURE):
        raw, encrypted, kind = _docx_text(data, deadline), False, "docx"
    else:
        raise KnowledgeFileError(KNOWLEDGE_FILE_UNSUPPORTED)
    text = _storable(raw).strip()
    if not text:
        raise _unreadable()
    return ExtractedDocument(
        kind=kind, text=text, size_bytes=len(data), encrypted=encrypted
    )


def _extract_all(files: Sequence[bytes], deadline: float) -> list[ExtractedDocument]:
    extracted: list[ExtractedDocument] = []
    for data in files:
        _check_deadline(deadline)
        extracted.append(extract_document_text(data, deadline=deadline))
    return extracted


async def extract_documents(
    files: Sequence[bytes], *, seconds: float = KNOWLEDGE_EXTRACTION_SECONDS
) -> list[ExtractedDocument]:
    """The text of every file of one upload, in a worker thread (never on
    the event loop), all of them within ``seconds``; the first file that
    cannot be ingested raises its :class:`KnowledgeFileError`."""

    deadline = time.monotonic() + seconds
    return await asyncio.to_thread(_extract_all, list(files), deadline)


class KnowledgeService:
    """
    Placeholder for Knowledge Service integration.

    When fully implemented, this will:
    1. Query the Qdrant vector database
    2. Retrieve relevant knowledge chunks
    3. Re-rank by relevance
    4. Return formatted chunks for context

    For now, returns empty results or mock data.
    """

    def __init__(
        self,
        qdrant_url: Optional[str] = None,
        collection_name: str = "knowledge",
        max_results: int = 5,
    ) -> None:
        """
        Initialize the Knowledge Service.

        Args:
            qdrant_url: URL to Qdrant server
            collection_name: Name of the collection
            max_results: Maximum results to return
        """
        self.qdrant_url = qdrant_url
        self.collection_name = collection_name
        self.max_results = max_results
        self._connected = False

    async def connect(self) -> bool:
        """
        Connect to the knowledge service.

        Returns:
            True if connected successfully
        """
        # Placeholder - would connect to Qdrant
        logger.info("Knowledge Service: Connection placeholder (not implemented)")
        self._connected = False
        return False

    async def disconnect(self) -> None:
        """Disconnect from the knowledge service."""
        self._connected = False

    async def query(
        self,
        tenant_id: UUID,
        query: str,
        max_results: Optional[int] = None,
        filters: Optional[dict[str, object]] = None,
    ) -> list[KnowledgeChunk]:
        """
        Query the knowledge base for relevant chunks.

        Args:
            tenant_id: Tenant ID to filter by
            query: Query string
            max_results: Maximum results to return
            filters: Additional filters

        Returns:
            List of relevant KnowledgeChunk objects
        """
        if not self._connected:
            logger.debug("Knowledge Service not connected, returning empty results")
            return []

        # Placeholder - would query Qdrant
        # Example of what this would do:
        # 1. Embed the query using the embedding model
        # 2. Query Qdrant with the embedding
        # 3. Filter by tenant_id
        # 4. Re-rank if needed
        # 5. Return top results

        return []

    async def query_by_intent(
        self,
        tenant_id: UUID,
        intent: str,
        query: str,
        max_results: Optional[int] = None,
    ) -> list[KnowledgeChunk]:
        """
        Query knowledge base by intent.

        Some intents may have specific knowledge associated.

        Args:
            tenant_id: Tenant ID
            intent: Detected intent
            query: Original query
            max_results: Maximum results

        Returns:
            List of relevant chunks
        """
        # Map intent to potential knowledge categories
        intent_categories = {
            "question_price": ["pricing", "values", "costs"],
            "question_location": ["location", "address", "directions"],
            "question_hours": ["hours", "schedule", "availability"],
            "question_services": ["services", "procedures", "treatments"],
            "question_insurance": ["insurance", "plans", "coverage"],
        }

        categories = intent_categories.get(intent, [])

        # Would filter by categories in actual implementation
        return await self.query(
            tenant_id=tenant_id,
            query=query,
            max_results=max_results,
            filters={"categories": categories} if categories else None,
        )

    async def get_chunk_by_id(
        self,
        tenant_id: UUID,
        chunk_id: str,
    ) -> Optional[KnowledgeChunk]:
        """
        Get a specific knowledge chunk by ID.

        Args:
            tenant_id: Tenant ID
            chunk_id: Chunk ID

        Returns:
            KnowledgeChunk if found
        """
        if not self._connected:
            return None

        # Placeholder
        return None

    def is_available(self) -> bool:
        """Check if knowledge service is available."""
        return self._connected

    async def health_check(self) -> dict[str, object]:
        """Perform health check on the knowledge service."""
        return {
            "connected": self._connected,
            "qdrant_url": self.qdrant_url,
            "collection": self.collection_name,
            "status": "not_implemented",
        }


# ============================================
# Mock Knowledge Service for Testing
# ============================================


class MockKnowledgeService(KnowledgeService):
    """
    Mock Knowledge Service for testing.

    Returns predefined knowledge chunks based on query keywords.
    """

    def __init__(self) -> None:
        super().__init__()
        self._connected = True
        self._mock_data = self._create_mock_data()

    def _create_mock_data(self) -> dict[str, list[KnowledgeChunk]]:
        """Create mock knowledge data."""
        return {
            "horário": [
                KnowledgeChunk(
                    id="mock_hours_1",
                    content="A clínica funciona de segunda a sexta, das 8h às 18h. Aos sábados, das 8h às 12h.",
                    source="FAQ",
                    relevance_score=0.9,
                ),
            ],
            "preço": [
                KnowledgeChunk(
                    id="mock_price_1",
                    content="A consulta inicial custa R$ 200,00. Retornos dentro de 30 dias são gratuitos.",
                    source="Pricing",
                    relevance_score=0.85,
                ),
            ],
            "endereço": [
                KnowledgeChunk(
                    id="mock_location_1",
                    content="Estamos localizados na Rua Exemplo, 123 - Centro. Próximo à estação de metrô.",
                    source="Contact",
                    relevance_score=0.9,
                ),
            ],
            "convênio": [
                KnowledgeChunk(
                    id="mock_insurance_1",
                    content="Aceitamos os convênios: Unimed, Bradesco Saúde, SulAmérica e Amil.",
                    source="Insurance",
                    relevance_score=0.88,
                ),
            ],
        }

    async def query(
        self,
        tenant_id: UUID,
        query: str,
        max_results: Optional[int] = None,
        filters: Optional[dict[str, object]] = None,
    ) -> list[KnowledgeChunk]:
        """Return mock knowledge based on query keywords."""
        query_lower = query.lower()
        results: list[KnowledgeChunk] = []

        for keyword, chunks in self._mock_data.items():
            if keyword in query_lower:
                results.extend(chunks)

        # Limit results
        max_r = max_results or self.max_results
        return results[:max_r]


# ============================================
# Singleton Instance
# ============================================

_service: Optional[KnowledgeService] = None


def get_knowledge_service(mock: bool = False) -> KnowledgeService:
    """
    Get or create the global KnowledgeService instance.

    Args:
        mock: If True, return a mock service for testing
    """
    global _service

    if mock:
        return MockKnowledgeService()

    if _service is None:
        _service = KnowledgeService()

    return _service
