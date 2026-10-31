"""Regenerate the Group 4a parser fixtures of this directory (provenance).

Run once with the locked ``pypdf`` and ``python-docx``; the committed bytes
are what the tests read (never this script): ``python
tests/fixtures/credit_storage_conformance/generate.py <out-dir>``.  With
``pypdf`` 6.21.4 every PDF and the corrupt files come out byte-identical;
the DOCX archive carries the run's zip timestamps, so only its bytes differ.
See ``README.md`` for what each file must do.
"""

from __future__ import annotations

import io
import sys
from datetime import datetime, timezone
from pathlib import Path


def raw_pdf(content: bytes, resources: bytes, extra_objects: list[bytes]) -> bytes:
    """A one-page PDF whose cross-reference table is computed exactly."""

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources " + resources + b" >>",
        b"<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
        *extra_objects,
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def main(root: Path) -> None:
    from docx import Document
    from pypdf import PdfReader, PdfWriter

    def normalized(data: bytes) -> bytes:
        writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
        out = io.BytesIO()
        writer.write(out)
        return out.getvalue()

    (root / "positive").mkdir(parents=True, exist_ok=True)
    (root / "negative").mkdir(parents=True, exist_ok=True)

    text_stream = (
        b"BT /F1 12 Tf 72 720 Td (Horario de atendimento: segunda a sexta, "
        b"das 8h as 18h.) Tj 0 -16 Td (A consulta inicial custa R$ 200,00.) Tj ET"
    )
    textual = normalized(
        raw_pdf(
            text_stream,
            b"<< /Font << /F1 5 0 R >> >>",
            [b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"],
        )
    )
    (root / "positive" / "textual.pdf").write_bytes(textual)

    image = bytes(range(0, 256, 16)) * 4  # 8 x 8 gray, no text at all
    ocr_only = normalized(
        raw_pdf(
            b"q 200 0 0 200 72 500 cm /Im1 Do Q",
            b"<< /XObject << /Im1 5 0 R >> >>",
            [
                b"<< /Type /XObject /Subtype /Image /Width 8 /Height 8 "
                b"/ColorSpace /DeviceGray /BitsPerComponent 8 /Length "
                + str(len(image)).encode()
                + b" >>\nstream\n"
                + image
                + b"\nendstream"
            ],
        )
    )
    (root / "negative" / "ocr-only.pdf").write_bytes(ocr_only)

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(textual)))
    writer.encrypt(
        user_password="segredo-do-teste",
        owner_password="dono-do-teste",
        algorithm="RC4-128",
    )
    protected = io.BytesIO()
    writer.write(protected)
    (root / "negative" / "password-protected.pdf").write_bytes(protected.getvalue())

    (root / "negative" / "corrupt.pdf").write_bytes(
        b"%PDF-1.7\n1 0 obj << /Type /Catalog /Pages 2 0 R\n\x00\xff\xfe truncated"
    )

    document = Document()
    document.add_heading("Clinica NeuroStack", level=1)
    document.add_paragraph("Atendemos de segunda a sexta, das 8h as 18h.")
    document.add_paragraph("Aceitamos os convenios Unimed e SulAmerica.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Consulta inicial"
    table.rows[0].cells[1].text = "R$ 200,00"
    fixed = datetime(2026, 10, 30, tzinfo=timezone.utc)
    document.core_properties.created = fixed
    document.core_properties.modified = fixed
    document.core_properties.author = "NeuroStack tests"
    out = io.BytesIO()
    document.save(out)
    (root / "positive" / "textual.docx").write_bytes(out.getvalue())

    (root / "negative" / "corrupt.docx").write_bytes(
        b"PK\x03\x04\x14\x00\x00\x00 this is not a zip archive \x00\xff"
    )

    # Group 5: an invalid format (a PNG signature: neither PDF nor DOCX).
    (root / "negative" / "invalid-format.bin").write_bytes(
        b"\x89PNG\r\n\x1a\nthis is neither a PDF nor a DOCX"
    )


if __name__ == "__main__":
    main(Path(sys.argv[1]))
